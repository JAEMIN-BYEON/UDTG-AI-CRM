#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
견적서 PPTX → 원문 텍스트 덤프 (표/도형/노트 전부)
에이전트가 이 출력을 읽고 24개 항목 JSON으로 매핑한다.

사용법: python3 parse_pptx.py "견적서.pptx"
"""
import subprocess
import sys


def ensure_pptx():
    try:
        import pptx  # noqa
    except ImportError:
        subprocess.run([sys.executable, "-m", "pip", "install", "python-pptx",
                        "--break-system-packages", "-q"], check=False)


def dump(path):
    ensure_pptx()
    from pptx import Presentation
    prs = Presentation(path)
    for i, slide in enumerate(prs.slides, 1):
        print("\n" + "=" * 60)
        print("[슬라이드 %d]" % i)
        print("=" * 60)
        for shape in slide.shapes:
            if shape.has_table:
                tbl = shape.table
                print("\n--- 표 (%d행 x %d열) ---" % (len(tbl.rows), len(tbl.columns)))
                for r in tbl.rows:
                    cells = [c.text.strip().replace("\n", " / ") for c in r.cells]
                    print(" | ".join(cells))
            elif shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    print("\n[TEXT] " + t)
            elif shape.shape_type == 13:  # PICTURE
                print("\n[IMAGE] %s" % (shape.name,))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            print("\n[NOTES] " + slide.notes_slide.notes_text_frame.text.strip())


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit("usage: parse_pptx.py <견적서.pptx>")
    dump(sys.argv[1])
